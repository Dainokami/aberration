import json
from pathlib import Path
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill, Border, Side, Alignment
from openpyxl.comments import Comment
from openpyxl.utils import get_column_letter

root = Path(__file__).resolve().parents[1]
manifest = json.loads((root / "assets/derived_art/scene_pack_20261016/manifest.json").read_text())
validation = json.loads((root / "docs/ocean_scene_assets_validation.json").read_text())
passed = {entry["id"]: entry["passed"] for entry in validation["assets"]}
names = {"ocean": "海洋", "grass": "草原", "desert": "沙漠", "rock": "岩石"}
wb = Workbook()
ws = wb.active
ws.title = "逐件验收"
headers = ["源目录", "源文件", "类别", "生态", "规格与视图", "设计表行", "资源ID",
           "原图尺寸(px)", "导入尺寸(px)", "画面尺寸(m)", "碰撞/触发", "格子中心偏移(m)",
           "单格区域(m)", "支撑要求", "可调机制参数", "配置检查", "视觉检查", "说明"]
ws.append(headers)
for entry in manifest["entries"]:
    kind = entry["kind"]
    note = ""
    if kind == "fruit_tree":
        note = "画面4格/3格视图；碰撞统一3格L形。384方图等比放大为512。"
    elif kind == "chest":
        note = "使用原图分区铰接开盖，掉落物采用物理抛物线。"
    elif kind == "bubbles":
        note = "仅海底；画面保持原始比例，水面处裁切。1×3为纵向画布。"
    if entry["parameters"]:
        note += " 原表未量化数值，所列为可调默认值。"
    ws.append([entry["source_directory"], entry["source_file"], entry["category"],
               "/".join(names[e] for e in entry["ecosystems"]),
               f'{entry["layout"][0]}×{entry["layout"][1]} ' + entry["view"],
               entry["spec_row"], entry["id"],
               "×".join(map(str, entry["original_pixels"])),
               "×".join(map(str, entry["output_pixels"])),
               "×".join(map(str, entry["image_size_m"])),
               "实体L形" if entry["solid_collision"] else "可进入触发区",
               json.dumps(entry["spatial_cells_m"]), "×".join(map(str, entry["cell_size_m"])),
               "真实海底实体，上方独立海水" if entry["requires_seabed"] else "匹配生态的实体顶面/侧壁",
               json.dumps(entry["parameters"], ensure_ascii=False) if entry["parameters"] else "—",
               "通过" if passed.get(entry["id"]) else "失败", "Metal画面已核对", note])
    ws.cell(ws.max_row, 2).comment = Comment("Source SHA-256: " + entry["source_sha256"], "验收")
    ws.cell(ws.max_row, 6).comment = Comment("Source: /Users/bytedance/Downloads/畸变_场景机制设计.xlsx / Sheet1", "验收")
    ws.cell(ws.max_row, 7).hyperlink = (root / "assets/generated" / entry["id"] / "definition.tres").as_uri()
for row in ws:
    for cell in row:
        cell.font = Font(name="Arial", size=11, color="263244")
        cell.alignment = Alignment(vertical="top", wrap_text=True)
        cell.border = Border(bottom=Side(style="thin", color="D9DEE7"))
        cell.fill = PatternFill("solid", fgColor="FFFFFF" if cell.row % 2 == 0 else "F7F9FC")
for cell in ws[1]:
    cell.fill = PatternFill("solid", fgColor="233448")
    cell.font = Font(name="Arial", size=11, bold=True, color="FFFFFF")
for row in range(2, ws.max_row + 1):
    ws.cell(row, 16).fill = PatternFill("solid", fgColor="E9F4EC")
    ws.cell(row, 16).font = Font(name="Arial", bold=True, color="28643A")
    if ws.cell(row, 15).value != "—":
        ws.cell(row, 15).fill = PatternFill("solid", fgColor="FFF5DA")
    ws.row_dimensions[row].height = 62
for index, width in enumerate([25, 42, 17, 18, 30, 12, 35, 18, 18, 18, 22, 38, 18, 32, 48, 15, 23, 55], 1):
    ws.column_dimensions[get_column_letter(index)].width = width
ws.row_dimensions[1].height = 30
ws.freeze_panes = "C2"
ws.auto_filter.ref = ws.dimensions
ws.sheet_view.showGridLines = False
ws.sheet_properties.pageSetUpPr.fitToPage = True
ws.page_setup.orientation = "landscape"
ws.page_setup.paperSize = ws.PAPERSIZE_A3
ws.page_setup.fitToWidth = 1
ws.page_setup.fitToHeight = 0
ws.print_title_rows = "1:1"
output = root.parent / "海洋地块与场景资产_逐件验收.xlsx"
wb.save(output)
verified = load_workbook(output)
assert verified.active.max_row == 32
assert all(verified.active.cell(row, 16).value == "通过" for row in range(2, 33))
assert not any(cell.data_type in ("f", "e") for row in verified.active for cell in row)
print(output)
